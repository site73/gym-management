# -*- coding: utf-8 -*-
import re
import os
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

BASE = r'E:\WorkBuddy\健身房管理业务\10-展示汇报\健身房运营管理系统答辩'
SRC = os.path.join(BASE, '展讲话稿.md')
DST = os.path.join(BASE, '展讲话稿.docx')


def set_font(style, east, asc, size, bold=False):
    style.font.name = asc
    style.font.size = Pt(size)
    style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = rpr.makeelement(qn('w:rFonts'), {})
        rpr.append(rf)
    rf.set(qn('w:eastAsia'), east)


def add_row(table, cells):
    ncols = len(table.columns)
    if len(cells) > ncols:
        cells = cells[:ncols - 1] + [' '.join(cells[ncols - 1:])]
    cells = list(cells) + [''] * (ncols - len(cells))
    row = table.add_row().cells
    for i in range(ncols):
        row[i].text = cells[i]


def add_runs(par, text):
    for p in re.split(r'(\*\*.+?\*\*)', text):
        if not p:
            continue
        if p.startswith('**') and p.endswith('**'):
            r = par.add_run(p[2:-2])
            r.bold = True
        else:
            par.add_run(p)


doc = Document()
for sec in doc.sections:
    sec.top_margin = Cm(2.0); sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.0); sec.right_margin = Cm(2.0)
doc.core_properties.author = 'site73'
doc.core_properties.last_modified_by = 'site73'
set_font(doc.styles['Normal'], '宋体', 'Times New Roman', 11)
doc.styles['Normal'].paragraph_format.line_spacing = 1.4
set_font(doc.styles['Title'], '黑体', 'Times New Roman', 20, bold=True)
set_font(doc.styles['Heading 1'], '黑体', 'Times New Roman', 16, bold=True)
set_font(doc.styles['Heading 2'], '黑体', 'Times New Roman', 13.5, bold=True)
set_font(doc.styles['Heading 3'], '黑体', 'Times New Roman', 12, bold=True)

in_table = False
table = None
with open(SRC, encoding='utf-8') as f:
    for line in f.read().splitlines():
        s = line.rstrip()
        if not s.strip() or s.strip() == '---':
            continue
        if s.startswith('# '):
            p = doc.add_paragraph(style='Title')
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_runs(p, s[2:].strip())
        elif s.startswith('## '):
            p = doc.add_heading(level=1)
            add_runs(p, s[3:].strip())
        elif s.startswith('### '):
            p = doc.add_heading(level=2)
            add_runs(p, s[4:].strip())
        elif s.lstrip().startswith('|'):
            parts = [x.strip() for x in s.strip().strip('|').split('|')]
            if all(set(c) <= {'-', ':', ' '} for c in parts):
                continue
            if not in_table:
                table = doc.add_table(rows=1, cols=len(parts))
                table.style = 'Table Grid'
                add_row(table, parts)
                in_table = True
            else:
                add_row(table, parts)
        else:
            if in_table:
                in_table = False
                table = None
            if s.lstrip().startswith('> '):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.6)
                p.paragraph_format.space_after = Pt(6)
                add_runs(p, s.lstrip()[2:])
            elif s.lstrip().startswith('- [ ]'):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.74)
                p.add_run('☐ ' + s.lstrip()[5:])
            elif s.lstrip().startswith('- '):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.74)
                p.add_run('• ' + s.lstrip()[2:])
            else:
                p = doc.add_paragraph()
                add_runs(p, s.strip())

for t in doc.tables:
    for r in t.rows:
        for c in r.cells:
            for pp in c.paragraphs:
                for rr in pp.runs:
                    rr.font.name = 'Times New Roman'
                    rr.font.size = Pt(9)
                    rpr = rr._element.get_or_add_rPr()
                    rf = rpr.find(qn('w:rFonts'))
                    if rf is None:
                        rf = rpr.makeelement(qn('w:rFonts'), {})
                        rpr.append(rf)
                    rf.set(qn('w:eastAsia'), '宋体')

doc.save(DST)
print('saved', DST)
